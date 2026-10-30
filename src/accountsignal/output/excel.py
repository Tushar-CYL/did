"""Power BI-ready Excel export: flat, normalized sheets with typed columns and a
run_date, so Power BI 'Get Data -> Excel' models them with no manual cleanup.
No merged cells, no internal rank columns."""
from __future__ import annotations

import logging
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import config
from ..store import io

log = logging.getLogger("accountsignal.output.excel")

CATEGORIES = ("hiring", "vision", "partnerships", "financial", "expansion")
_HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb: Workbook, title: str, headers: list[str]) -> "Worksheet":
    ws = wb.create_sheet(title)
    ws.append(headers)
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    return ws


def _autosize(ws, max_width: int = 60) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, max_width)


def build(out_path: Path | None = None) -> Path:
    data = io.read_json(io.normalized_path())
    run_date = date.today().isoformat()
    wb = Workbook()
    wb.remove(wb.active)  # drop default sheet

    # 1) Companies overview
    ws = _sheet(wb, "Companies", [
        "run_date", "company", "sector", "ats", "hiring_india_count", "news_count",
        "insight", "summary_vision", "summary_partnerships", "summary_financial",
        "summary_expansion", "website", "career_url",
    ])
    for c in data:
        mi = c["market_intel"]
        news_count = sum(len(mi[k]["items"]) for k in CATEGORIES)
        ws.append([
            run_date, c["company"], c.get("sector") or "", c.get("ats") or "",
            len(c["hiring_india"]), news_count,
            " | ".join(c.get("signals", {}).get("insight", [])),
            mi["vision"]["summary"], mi["partnerships"]["summary"],
            mi["financial"]["summary"], mi["expansion"]["summary"],
            c.get("website") or "", c.get("career_url") or "",
        ])

    # 2) Hiring (India) — tools/aging centric, no critical/type/manager
    ws = _sheet(wb, "Hiring_India", [
        "run_date", "company", "sector", "role", "department", "team", "seniority",
        "skills_tools", "city", "region", "posted_date", "aging_days", "source_url",
    ])
    for c in data:
        for r in c["hiring_india"]:
            ws.append([
                run_date, c["company"], c.get("sector") or "", r.get("role", ""),
                r.get("department", ""), r.get("team", ""), r.get("seniority", ""),
                ", ".join(r.get("tools", [])), r.get("city", ""), r.get("region", ""),
                r.get("posted_date", ""), r.get("days_open"), r.get("source_url", ""),
            ])

    # 3) Skills / Tools (one row per tool occurrence) — the intel, for Power BI pivots
    ws = _sheet(wb, "Skills_Tools", ["run_date", "company", "sector", "department", "tool"])
    for c in data:
        for r in c["hiring_india"]:
            for t in r.get("tools", []):
                ws.append([run_date, c["company"], c.get("sector") or "", r.get("department", ""), t])

    # 4) News / market intelligence (NO internal score)
    ws = _sheet(wb, "News", ["run_date", "company", "category", "title", "date", "source", "url"])
    for c in data:
        mi = c["market_intel"]
        for cat in CATEGORIES:
            for it in mi[cat]["items"]:
                ws.append([run_date, c["company"], cat, it.get("title", ""),
                           it.get("date", ""), it.get("source", ""), it.get("url", "")])

    # 5) Signals / insight (long form)
    ws = _sheet(wb, "Signals", ["run_date", "company", "signal_type", "signal"])
    for c in data:
        sig = c.get("signals", {})
        for stype, key in [("Insight", "insight"), ("Partnership", "partnership_signals"),
                           ("Vision / expansion", "vision_shifts"), ("Tool gravitation", "tool_gravitation")]:
            for s in sig.get(key, []):
                ws.append([run_date, c["company"], stype, s])

    for ws in wb.worksheets:
        _autosize(ws)

    out_path = out_path or (config.OUT_DIR / f"AccountSignal_{run_date.replace('-', '')}.xlsx")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    log.info("Excel written -> %s", out_path)
    return out_path
